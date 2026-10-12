#!/usr/bin/env python3
"""Build report/REPORT.docx from report/REPORT.md (needs pandoc + python-docx).

  python3 scripts/build_report_docx.py
Pandoc's default tables have no borders; this adds borders, a header shade,
a smaller table font and narrower margins so the wide tables fit on A4.
"""
import os
import subprocess

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD = os.path.join(ROOT, "report", "REPORT.md")
OUT = os.path.join(ROOT, "report", "REPORT.docx")

subprocess.run(["pandoc", MD, "-o", OUT, "--resource-path", f"{ROOT}/report:{ROOT}"], check=True)

doc = Document(OUT)
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.8)

for table in doc.tables:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tbl_pr.append(borders)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if r_idx == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "E8EEF7")
                cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)
doc.save(OUT)
print("written", os.path.relpath(OUT, ROOT))
